from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from pdf2image import pdfinfo_from_path
from PIL import Image


docx_path = Path(r"D:\Desktop\OpsSentinel\张峥_求职简历_科研版.docx")
pdf_path = Path(r"D:\Desktop\OpsSentinel\.resume_qa\research\resume.pdf")
png_path = Path(r"D:\Desktop\OpsSentinel\.resume_qa\research\page-1.png")
poppler = r"C:\Users\张峥\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin"

assert docx_path.exists() and docx_path.stat().st_size > 20_000
doc = Document(docx_path)
assert len(doc.sections) == 1
section = doc.sections[0]
assert abs(section.page_width.cm - 21.0) < 0.05
assert abs(section.page_height.cm - 29.7) < 0.05

parts = [p.text for p in doc.paragraphs]
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            parts.extend(p.text for p in cell.paragraphs)
text = "\n".join(parts)

for required in (
    "大模型应用 / Agent 开发工程师",
    "Gist-Residual",
    "Oracle 准确率为 38%",
    "gh-assistant",
    "OpsSentinel",
    "Linux 全套测试 132 项通过",
    "CET-6 531",
):
    assert required in text, required
for forbidden in ("gh-assigh-assistant", "TODO", "TBD", "placeholder"):
    assert forbidden not in text, forbidden

with ZipFile(docx_path) as archive:
    names = set(archive.namelist())
    document_xml = archive.read("word/document.xml").decode("utf-8")
    styles_xml = archive.read("word/styles.xml").decode("utf-8")
    assert "txbxContent" not in document_xml
    assert re.search(r"<w:ins(?:\s|>)", document_xml) is None
    assert re.search(r"<w:del(?:\s|>)", document_xml) is None
    assert "word/comments.xml" not in names
    assert "Microsoft YaHei" in styles_xml
    assert "Arial" in styles_xml
    assert any(name.startswith("word/media/image") for name in names)

info = pdfinfo_from_path(str(pdf_path), poppler_path=poppler)
assert info["Pages"] == 1
with Image.open(png_path) as image:
    assert image.width >= 1200 and image.height >= 1700

print("DOCX_OK", docx_path.stat().st_size)
print("A4_OK", round(section.page_width.cm, 2), round(section.page_height.cm, 2))
print("CONTENT_OK", len(text), "characters")
print("STRUCTURE_OK", "textboxes=0", "tracked_changes=0", "comments=0")
print("FONTS_OK", "Microsoft YaHei", "Arial")
print("RENDER_OK", info["Pages"], "page", Image.open(png_path).size)
